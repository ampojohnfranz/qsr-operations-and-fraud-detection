import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from datetime import datetime

df = pd.read_csv('/mnt/user-data/outputs/pos_transaction_logs_v2.csv', parse_dates=['Timestamp'])

REFUND_THRESHOLD  = 30.00
VOID_LIMIT        = 5
DISCOUNT_LIMIT    = 5
LATE_NIGHT_START  = 22
LATE_NIGHT_END    = 6

df['Date'] = df['Timestamp'].dt.date
df['Hour'] = df['Timestamp'].dt.hour

flags = []

# 1. High-value refunds
high_refunds = df[(df['Transaction_Type'] == 'Refund') & (df['Amount_USD'] > REFUND_THRESHOLD)].copy()
high_refunds['Flag_Reason'] = f'High-value refund (>${REFUND_THRESHOLD})'
flags.append(high_refunds)

# 2. Excessive voids per cashier per day
void_counts = df[df['Transaction_Type'] == 'Void'].groupby(['Cashier_ID','Date']).size().reset_index(name='cnt')
excess_voids = void_counts[void_counts['cnt'] > VOID_LIMIT][['Cashier_ID','Date']]
void_flagged = df.merge(excess_voids, on=['Cashier_ID','Date'])
void_flagged = void_flagged[void_flagged['Transaction_Type'] == 'Void'].copy()
void_flagged['Flag_Reason'] = f'Excessive voids (>{VOID_LIMIT}/day by same cashier)'
flags.append(void_flagged)

# 3. Excessive discounts per cashier per day
disc_counts = df[df['Transaction_Type'] == 'Discount applied'].groupby(['Cashier_ID','Date']).size().reset_index(name='cnt')
excess_disc = disc_counts[disc_counts['cnt'] > DISCOUNT_LIMIT][['Cashier_ID','Date']]
disc_flagged = df.merge(excess_disc, on=['Cashier_ID','Date'])
disc_flagged = disc_flagged[disc_flagged['Transaction_Type'] == 'Discount applied'].copy()
disc_flagged['Flag_Reason'] = f'Excessive discounts (>{DISCOUNT_LIMIT}/day by same cashier)'
flags.append(disc_flagged)

# 4. Late-night voids or refunds
late = df[
    ((df['Hour'] >= LATE_NIGHT_START) | (df['Hour'] < LATE_NIGHT_END)) &
    (df['Transaction_Type'].isin(['Void','Refund']))
].copy()
late['Flag_Reason'] = 'Late-night void/refund (10PM–6AM)'
flags.append(late)

flagged = pd.concat(flags).drop_duplicates(subset='Transaction_ID', keep='first')
flagged = flagged.sort_values(['Store','Cashier_ID','Timestamp'])

flag_rate = round(len(flagged) / len(df) * 100, 1)

summary_store = flagged.groupby('Store').agg(
    Total_Flags=('Transaction_ID','count'),
    Total_Amount=('Amount_USD','sum')
).reset_index()

summary_cashier = flagged.groupby(['Cashier_ID','Store']).agg(
    Total_Flags=('Transaction_ID','count'),
    Flag_Types=('Flag_Reason', lambda x: ' | '.join(sorted(x.unique())))
).reset_index().sort_values('Total_Flags', ascending=False)

# ── Excel ──────────────────────────────────────────────────────────────────────
wb = Workbook()

RED_FILL    = PatternFill('solid', start_color='FFCCCC')
ORANGE_FILL = PatternFill('solid', start_color='FFE0B2')
YELLOW_FILL = PatternFill('solid', start_color='FFF9C4')
GREY_FILL   = PatternFill('solid', start_color='EEEEEE')
HEADER_FILL = PatternFill('solid', start_color='1F2D3D')
GREEN_FILL  = PatternFill('solid', start_color='C8E6C9')
TITLE_FILL  = PatternFill('solid', start_color='0D1B2A')

HEADER_FONT = Font(name='Arial', bold=True, color='FFFFFF', size=10)
TITLE_FONT  = Font(name='Arial', bold=True, size=16, color='FFFFFF')
META_FONT   = Font(name='Arial', size=9,  color='AAAAAA', italic=True)
SUB_FONT    = Font(name='Arial', bold=True, size=10, color='1F2D3D')
BODY_FONT   = Font(name='Arial', size=10)
CENTER      = Alignment(horizontal='center', vertical='center')
LEFT        = Alignment(horizontal='left',   vertical='center', wrap_text=False)
WRAP        = Alignment(horizontal='left',   vertical='center', wrap_text=True)
thin        = Side(style='thin', color='DDDDDD')
BORDER      = Border(left=thin, right=thin, top=thin, bottom=thin)

def set_header(ws, row, ncols):
    for c in range(1, ncols+1):
        cell = ws.cell(row=row, column=c)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = CENTER
        cell.border = BORDER

def set_body(ws, row, ncols, fill=None):
    for c in range(1, ncols+1):
        cell = ws.cell(row=row, column=c)
        if fill: cell.fill = fill
        cell.font = BODY_FONT
        cell.alignment = LEFT
        cell.border = BORDER

# ── Sheet 1: Flagged Transactions ──────────────────────────────────────────────
ws1 = wb.active
ws1.title = 'Flagged Transactions'

# Title banner
ws1.merge_cells('A1:H1')
ws1['A1'] = 'POS Suspicious Activity Report — May 2026'
ws1['A1'].font = TITLE_FONT
ws1['A1'].fill = TITLE_FILL
ws1['A1'].alignment = CENTER
ws1.row_dimensions[1].height = 36

ws1.merge_cells('A2:H2')
ws1['A2'] = (f'Generated: {datetime.now().strftime("%B %d, %Y")}   |   '
             f'Total transactions: {len(df):,}   |   '
             f'Flagged: {len(flagged):,} ({flag_rate}%)   |   '
             f'Source: pos_transaction_logs_v2.csv')
ws1['A2'].font = META_FONT
ws1['A2'].alignment = CENTER
ws1.row_dimensions[2].height = 18

headers = ['Transaction ID','Timestamp','Store','Cashier ID','Type','Payment Method','Amount (USD)','Flag Reason']
widths  = [15, 22, 10, 12, 20, 16, 14, 44]
for i, (h, w) in enumerate(zip(headers, widths), 1):
    ws1.column_dimensions[get_column_letter(i)].width = w

ws1.append(headers)
set_header(ws1, 3, len(headers))
ws1.row_dimensions[3].height = 20

data_cols = ['Transaction_ID','Timestamp','Store','Cashier_ID','Transaction_Type','Payment_Method','Amount_USD','Flag_Reason']
for i, row in enumerate(flagged[data_cols].values.tolist(), 4):
    reason = str(row[7])
    if 'High-value refund' in reason:    fill = RED_FILL
    elif 'Excessive voids' in reason:    fill = ORANGE_FILL
    elif 'Excessive discounts' in reason: fill = YELLOW_FILL
    else:                                fill = GREY_FILL

    for j, val in enumerate(row, 1):
        cell = ws1.cell(row=i, column=j, value=val)
        cell.fill = fill
        cell.font = BODY_FONT
        cell.border = BORDER
        cell.alignment = LEFT

ws1.freeze_panes = 'A4'
ws1.auto_filter.ref = f'A3:{get_column_letter(len(headers))}{3+len(flagged)}'

# ── Sheet 2: Summary by Store ──────────────────────────────────────────────────
ws2 = wb.create_sheet('Summary by Store')
ws2.merge_cells('A1:C1')
ws2['A1'] = 'Flag Summary by Store'
ws2['A1'].font = Font(name='Arial', bold=True, size=13, color='FFFFFF')
ws2['A1'].fill = TITLE_FILL
ws2['A1'].alignment = CENTER
ws2.row_dimensions[1].height = 30

for i, (h, w) in enumerate(zip(['Store','Total Flags','Amount at Risk (USD)'],[14,14,24]), 1):
    ws2.column_dimensions[get_column_letter(i)].width = w

ws2.append(['Store','Total Flags','Amount at Risk (USD)'])
set_header(ws2, 2, 3)

max_flags = summary_store['Total_Flags'].max()
for _, row in summary_store.iterrows():
    r = ws2.max_row + 1
    ws2.cell(row=r, column=1, value=row['Store'])
    ws2.cell(row=r, column=2, value=row['Total_Flags'])
    ws2.cell(row=r, column=3, value=round(row['Total_Amount'], 2))
    fill = RED_FILL if row['Total_Flags'] == max_flags else (ORANGE_FILL if row['Total_Flags'] >= max_flags*0.7 else GREY_FILL)
    set_body(ws2, r, 3, fill)

tr = ws2.max_row + 1
ws2.cell(row=tr, column=1, value='TOTAL').font = SUB_FONT
ws2.cell(row=tr, column=2, value=f'=SUM(B3:B{tr-1})').font = SUB_FONT
ws2.cell(row=tr, column=3, value=f'=SUM(C3:C{tr-1})').font = SUB_FONT
for c in range(1,4):
    ws2.cell(row=tr, column=c).fill = GREEN_FILL
    ws2.cell(row=tr, column=c).border = BORDER
    ws2.cell(row=tr, column=c).alignment = LEFT

# ── Sheet 3: Summary by Cashier ────────────────────────────────────────────────
ws3 = wb.create_sheet('Summary by Cashier')
ws3.merge_cells('A1:D1')
ws3['A1'] = 'Flag Summary by Cashier — Sorted by Risk'
ws3['A1'].font = Font(name='Arial', bold=True, size=13, color='FFFFFF')
ws3['A1'].fill = TITLE_FILL
ws3['A1'].alignment = CENTER
ws3.row_dimensions[1].height = 30

for i, (h, w) in enumerate(zip(['Cashier ID','Store','Total Flags','Flag Types'],[12,10,12,55]), 1):
    ws3.column_dimensions[get_column_letter(i)].width = w

ws3.append(['Cashier ID','Store','Total Flags','Flag Types'])
set_header(ws3, 2, 4)

max_c = summary_cashier['Total_Flags'].max()
for _, row in summary_cashier.iterrows():
    r = ws3.max_row + 1
    ws3.cell(row=r, column=1, value=row['Cashier_ID'])
    ws3.cell(row=r, column=2, value=row['Store'])
    ws3.cell(row=r, column=3, value=row['Total_Flags'])
    ws3.cell(row=r, column=4, value=row['Flag_Types'])
    fill = RED_FILL if row['Total_Flags'] == max_c else (ORANGE_FILL if row['Total_Flags'] >= max_c*0.6 else GREY_FILL)
    set_body(ws3, r, 4, fill)
    ws3.cell(row=r, column=4).alignment = WRAP
    ws3.row_dimensions[r].height = 30

# ── Sheet 4: Legend ────────────────────────────────────────────────────────────
ws4 = wb.create_sheet('Legend & Thresholds')
ws4.merge_cells('A1:B1')
ws4['A1'] = 'Report Legend & Detection Thresholds'
ws4['A1'].font = Font(name='Arial', bold=True, size=13, color='FFFFFF')
ws4['A1'].fill = TITLE_FILL
ws4['A1'].alignment = CENTER
ws4.row_dimensions[1].height = 30
ws4.column_dimensions['A'].width = 32
ws4.column_dimensions['B'].width = 52

sections = [
    ('COLOR CODING', None),
    ('Red',    f'High-value refund — amount exceeds ${REFUND_THRESHOLD}'),
    ('Orange', f'Excessive voids — cashier exceeded {VOID_LIMIT} voids in a single day'),
    ('Yellow', f'Excessive discounts — cashier exceeded {DISCOUNT_LIMIT} discounts in a single day'),
    ('Gray',   'Late-night void or refund — transaction between 10PM and 6AM'),
    ('', ''),
    ('DETECTION THRESHOLDS', None),
    ('Refund threshold',      f'${REFUND_THRESHOLD} USD'),
    ('Max voids/cashier/day', f'{VOID_LIMIT} transactions'),
    ('Max discounts/day',     f'{DISCOUNT_LIMIT} transactions'),
    ('Late-night window',     '10:00 PM – 6:00 AM'),
    ('', ''),
    ('REPORT STATS', None),
    ('Total transactions',  f'{len(df):,}'),
    ('Flagged transactions', f'{len(flagged):,} ({flag_rate}%)'),
    ('Date range',          'May 1–30, 2026'),
    ('Script',              'fraud_detector_v2.py'),
]

color_map = {'Red': RED_FILL, 'Orange': ORANGE_FILL, 'Yellow': YELLOW_FILL, 'Gray': GREY_FILL}
for r, (label, value) in enumerate(sections, 2):
    ws4.cell(row=r, column=1, value=label)
    ws4.cell(row=r, column=2, value=value if value else '')
    if value is None:
        ws4.cell(row=r, column=1).fill = HEADER_FILL
        ws4.cell(row=r, column=1).font = HEADER_FONT
        ws4.merge_cells(f'A{r}:B{r}')
        ws4.cell(row=r, column=1).alignment = CENTER
    elif label in color_map:
        ws4.cell(row=r, column=1).fill = color_map[label]
        ws4.cell(row=r, column=1).font = BODY_FONT
        ws4.cell(row=r, column=2).font = BODY_FONT
    else:
        ws4.cell(row=r, column=1).font = SUB_FONT
        ws4.cell(row=r, column=2).font = BODY_FONT

output = '/mnt/user-data/outputs/POS_Fraud_Detection_Report_v2.xlsx'
wb.save(output)
print(f"Done. {len(flagged)} flagged out of {len(df):,} ({flag_rate}%)")
print(f"Top flagged cashiers:\n{summary_cashier[['Cashier_ID','Store','Total_Flags']].head(5).to_string(index=False)}")
